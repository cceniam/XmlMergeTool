"""用于加载、勾选与合并 Robot 输出套件的应用服务。"""

from __future__ import annotations

import copy
import csv
import io
import math
import os
import re
import sys
import tempfile
import uuid
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, Sequence

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from robot.rebot import Rebot

from xml_merge_tool.application.models import (
    ExcelExportResult,
    LoadBatchResult,
    MergeResult,
    RebotOptions,
    SelectionSuggestion,
    SuiteCandidate,
)
from xml_merge_tool.application.selection_policy import SuiteSelectionPolicy
from xml_merge_tool.infrastructure.rf_reader import RFResultReader
from xml_merge_tool.infrastructure.robot_xml import parse_robot_tree


class XmlMergeService:
    """供 GUI/控制器调用的高层服务。"""

    def __init__(
        self,
        reader: RFResultReader | None = None,
        policy: SuiteSelectionPolicy | None = None,
    ):
        self._reader = reader or RFResultReader()
        self._policy = policy or SuiteSelectionPolicy()

    def load_candidates(self, xml_paths: Sequence[str]) -> LoadBatchResult:
        normalized_paths = self._normalize_paths(xml_paths)
        if not normalized_paths:
            return LoadBatchResult(
                candidates=tuple(),
                issues=tuple(),
                parsed_file_count=0,
                failed_file_count=0,
                successful_file_paths=tuple(),
            )
        return self._reader.load_candidates(normalized_paths)

    def suggest_selection(
        self, candidates: Iterable[SuiteCandidate]
    ) -> SelectionSuggestion:
        return self._policy.suggest(candidates)

    def merge_selected(
        self,
        candidates: Sequence[SuiteCandidate],
        selected_ids: Sequence[str],
        output_dir: str,
        rebot_options: RebotOptions | None = None,
    ) -> MergeResult:
        selected_set = {item for item in selected_ids if item}
        if not selected_set:
            raise ValueError("请至少勾选一个测试套件后再合并。")

        candidate_by_id = {candidate.candidate_id: candidate for candidate in candidates}
        selected_candidates = [
            candidate_by_id[candidate_id]
            for candidate_id in selected_ids
            if candidate_id in candidate_by_id
        ]

        if not selected_candidates:
            raise ValueError("勾选项无效，请重新加载 XML 后再试。")

        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")

        merged_xml_path = os.path.join(output_dir, f"merged-{timestamp}.xml")
        conflict_report_path = os.path.join(
            output_dir, f"conflict_report-{timestamp}.csv"
        )
        log_path = os.path.join(output_dir, f"log-{timestamp}.html")
        report_path = os.path.join(output_dir, f"report-{timestamp}.html")

        self._write_merged_xml(selected_candidates, merged_xml_path)
        self._write_conflict_report(
            candidates,
            selected_set,
            self.suggest_selection(candidates),
            conflict_report_path,
        )
        self._generate_robot_report(
            merged_xml_path,
            log_path,
            report_path,
            output_dir,
            rebot_options or RebotOptions(),
        )

        return MergeResult(
            merged_xml_path=merged_xml_path,
            log_path=log_path,
            report_path=report_path,
            conflict_report_path=conflict_report_path,
            selected_count=len(selected_candidates),
            candidate_count=len(candidates),
        )

    def generate_single_candidate_log(
        self,
        candidate: SuiteCandidate,
        output_dir: str,
        rebot_options: RebotOptions | None = None,
        *,
        log_filename: str = "single-log.html",
    ) -> str:
        if not candidate:
            raise ValueError("请选择有效的测试记录。")
        if not os.path.exists(candidate.file_path):
            raise ValueError(f"源文件不存在: {candidate.file_path}")

        os.makedirs(output_dir, exist_ok=True)
        log_path = os.path.join(output_dir, log_filename)
        temp_xml_path = os.path.join(
            output_dir, f".single-log-input-{uuid.uuid4().hex}.xml"
        )

        try:
            self._write_merged_xml([candidate], temp_xml_path)
            self._generate_robot_log(
                temp_xml_path,
                log_path,
                output_dir,
                rebot_options or RebotOptions(),
            )
            return log_path
        finally:
            try:
                if os.path.exists(temp_xml_path):
                    os.remove(temp_xml_path)
            except OSError:
                pass

    def export_single_candidate_xml(
        self,
        candidate: SuiteCandidate,
        output_path: str,
    ) -> str:
        """将单个候选套件导出为可供 Robot/Rebot 使用的独立 XML。"""
        if not candidate:
            raise ValueError("请选择有效的测试记录。")
        if not os.path.exists(candidate.file_path):
            raise ValueError(f"源文件不存在: {candidate.file_path}")
        if not output_path:
            raise ValueError("请选择有效的导出路径。")

        resolved_output_path = os.path.abspath(output_path)
        if os.path.normcase(resolved_output_path) == os.path.normcase(
            os.path.abspath(candidate.file_path)
        ):
            raise ValueError("导出路径不能覆盖源 XML 文件。")

        output_dir = os.path.dirname(resolved_output_path)
        if not output_dir:
            raise ValueError("请选择有效的导出目录。")
        os.makedirs(output_dir, exist_ok=True)
        self._write_merged_xml([candidate], resolved_output_path)
        return resolved_output_path

    def export_selected_tests_excel(
        self,
        candidates: Sequence[SuiteCandidate],
        selected_ids: Sequence[str],
        output_path: str,
    ) -> ExcelExportResult:
        """将全部已勾选套件中的直接测试结果导出为单个 Excel 工作表。"""
        selected_candidates = self._resolve_selected_candidates(candidates, selected_ids)
        rows: list[tuple[str, str, str, float | None, str]] = []
        for candidate in selected_candidates:
            robot_file_name = self._resolve_robot_file_name(candidate)
            for test in candidate.test_results:
                status = self._normalize_test_status(test.status)
                rows.append(
                    (
                        test.name,
                        robot_file_name,
                        status,
                        test.duration_seconds,
                        test.failure_reason if status == "FAIL" else "",
                    )
                )

        if not rows:
            raise ValueError("勾选套件中没有可导出的测试结果，请重新加载 XML 后再试。")
        if len(rows) > 1_048_575:
            raise ValueError("勾选测试数量超过单个 Excel 工作表上限。")
        if not output_path:
            raise ValueError("请选择有效的 Excel 导出路径。")

        resolved_output_path = os.path.abspath(output_path)
        if Path(resolved_output_path).suffix.lower() != ".xlsx":
            raise ValueError("Excel 导出路径必须使用 .xlsx 扩展名。")
        output_dir = os.path.dirname(resolved_output_path)
        if not output_dir:
            raise ValueError("请选择有效的 Excel 导出目录。")

        temp_path: str | None = None
        try:
            os.makedirs(output_dir, exist_ok=True)
            temp_fd, temp_path = tempfile.mkstemp(
                dir=output_dir,
                prefix=f".{Path(resolved_output_path).stem}-",
                suffix=".xlsx",
            )
            os.close(temp_fd)
            self._write_test_results_workbook(rows, temp_path)
            self._verify_test_results_workbook(temp_path, expected_rows=len(rows))
            os.replace(temp_path, resolved_output_path)
        except PermissionError as exc:
            raise ValueError(
                "Excel 文件无法写入，请关闭已打开的目标文件或选择其他位置。"
            ) from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Excel 导出失败: {exc}") from exc
        finally:
            try:
                if temp_path and os.path.exists(temp_path):
                    os.remove(temp_path)
            except OSError:
                pass

        return ExcelExportResult(
            output_path=resolved_output_path,
            selected_suite_count=len(selected_candidates),
            test_count=len(rows),
        )

    @staticmethod
    def _resolve_selected_candidates(
        candidates: Sequence[SuiteCandidate], selected_ids: Sequence[str]
    ) -> list[SuiteCandidate]:
        normalized_ids: list[str] = []
        seen: set[str] = set()
        for candidate_id in selected_ids:
            if not candidate_id or candidate_id in seen:
                continue
            seen.add(candidate_id)
            normalized_ids.append(candidate_id)
        if not normalized_ids:
            raise ValueError("请至少勾选一个测试套件。")

        candidate_by_id = {candidate.candidate_id: candidate for candidate in candidates}
        invalid_ids = [
            candidate_id
            for candidate_id in normalized_ids
            if candidate_id not in candidate_by_id
        ]
        if invalid_ids:
            raise ValueError("勾选项已失效，请重新加载 XML 文件后再试。")

        selected_set = set(normalized_ids)
        return [
            candidate
            for candidate in candidates
            if candidate.candidate_id in selected_set
        ]

    @staticmethod
    def _resolve_robot_file_name(candidate: SuiteCandidate) -> str:
        source = (candidate.robot_source or "").strip().rstrip("/\\")
        if source:
            basename = re.split(r"[\\/]", source)[-1].strip()
            if basename:
                return basename
        return candidate.file_name

    @staticmethod
    def _normalize_test_status(status: str | None) -> str:
        normalized = (status or "").strip().upper()
        if normalized in {"PASS", "FAIL", "SKIP"}:
            return normalized
        return "UNKNOWN"

    @classmethod
    def _write_test_results_workbook(
        cls,
        rows: Sequence[tuple[str, str, str, float | None, str]],
        output_path: str,
    ) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "测试结果"
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_view.zoomScale = 90
        worksheet.sheet_properties.tabColor = "1F4E78"
        worksheet.freeze_panes = "A2"

        headers = ("测试名称", "Robot 文件名", "状态", "耗时", "失败原因")
        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(
            name="Microsoft YaHei", size=11, color="FFFFFF", bold=True
        )
        body_font = Font(name="Microsoft YaHei", size=10.5, color="1F2937")
        banded_fills = (
            PatternFill(fill_type="solid", fgColor="FFFFFF"),
            PatternFill(fill_type="solid", fgColor="EAF2F8"),
        )
        header_border = Border(
            bottom=Side(style="medium", color="163A5C")
        )
        body_border = Border(
            bottom=Side(style="thin", color="D6E2EC")
        )
        header_alignment = Alignment(horizontal="center", vertical="center")
        text_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        status_alignment = Alignment(horizontal="center", vertical="center")

        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=column, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = header_border
        worksheet.row_dimensions[1].height = 26

        for row_index, values in enumerate(rows, start=2):
            test_name, robot_file_name, status, duration_seconds, failure_reason = values
            safe_values = (
                cls._excel_text(test_name),
                cls._excel_text(robot_file_name),
                cls._excel_text(status),
                cls._excel_text(failure_reason),
            )
            text_values = {
                1: safe_values[0],
                2: safe_values[1],
                3: safe_values[2],
                5: safe_values[3],
            }
            for column, value in text_values.items():
                cell = worksheet.cell(row=row_index, column=column)
                cell.value = value
                cell.data_type = "s"
                cell.font = body_font
                cell.alignment = status_alignment if column == 3 else text_alignment

            duration_cell = worksheet.cell(row=row_index, column=4)
            duration_cell.value = cls._excel_duration_value(duration_seconds)
            duration_cell.number_format = "[h]:mm:ss.000"
            duration_cell.font = body_font
            duration_cell.alignment = status_alignment

            row_fill = banded_fills[(row_index - 2) % len(banded_fills)]
            for cell in worksheet[row_index]:
                cell.fill = row_fill
                cell.border = body_border

            estimated_lines = max(
                1,
                (len(safe_values[0]) + 39) // 40,
                (len(safe_values[3]) + 69) // 70,
                safe_values[0].count("\n") + 1,
                safe_values[3].count("\n") + 1,
            )
            worksheet.row_dimensions[row_index].height = min(90, max(22, estimated_lines * 18))

        last_row = len(rows) + 1
        worksheet.auto_filter.ref = f"A1:E{last_row}"
        worksheet.column_dimensions["A"].width = 44
        worksheet.column_dimensions["B"].width = 36
        worksheet.column_dimensions["C"].width = 12
        worksheet.column_dimensions["D"].width = 18
        worksheet.column_dimensions["E"].width = 80

        status_styles = {
            "PASS": ("DFF5E9", "0F7A45"),
            "SKIP": ("FFF2DD", "9A5D00"),
            "FAIL": ("FDE7E7", "B42318"),
        }
        status_range = f"C2:C{last_row}"
        for status, (fill_color, font_color) in status_styles.items():
            worksheet.conditional_formatting.add(
                status_range,
                FormulaRule(
                    formula=[f'$C2="{status}"'],
                    fill=PatternFill(fill_type="solid", fgColor=fill_color),
                    font=Font(name="Microsoft YaHei", color=font_color, bold=True),
                ),
            )

        workbook.save(output_path)
        workbook.close()

    @staticmethod
    def _verify_test_results_workbook(output_path: str, *, expected_rows: int) -> None:
        workbook = load_workbook(output_path, read_only=False, data_only=False)
        try:
            if workbook.sheetnames != ["测试结果"]:
                raise ValueError("Excel 工作表结构验证失败。")
            worksheet = workbook["测试结果"]
            headers = tuple(
                worksheet.cell(row=1, column=column).value
                for column in range(1, 6)
            )
            if headers != ("测试名称", "Robot 文件名", "状态", "耗时", "失败原因"):
                raise ValueError("Excel 表头验证失败。")
            if worksheet.max_row != expected_rows + 1 or worksheet.max_column != 5:
                raise ValueError("Excel 数据行数验证失败。")
            for row in range(2, expected_rows + 2):
                if (
                    worksheet.cell(row=row, column=4).number_format
                    != "[h]:mm:ss.000"
                ):
                    raise ValueError("Excel 耗时列格式验证失败。")
        finally:
            workbook.close()

    @staticmethod
    def _excel_duration_value(duration_seconds: float | None) -> float | None:
        if duration_seconds is None:
            return None
        seconds = float(duration_seconds)
        if not math.isfinite(seconds) or seconds < 0:
            return None
        return seconds / 86_400

    @staticmethod
    def _excel_text(value: object) -> str:
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "�", str(value or ""))
        limit = 32_767
        marker = "\n[内容过长，已截断]"
        if len(text) <= limit:
            return text
        return text[: limit - len(marker)] + marker

    @staticmethod
    def _normalize_paths(paths: Sequence[str]) -> list[str]:
        resolved: list[str] = []
        seen: set[str] = set()
        for path in paths:
            if not path:
                continue
            abs_path = os.path.abspath(path)
            if abs_path in seen or not abs_path.lower().endswith(".xml"):
                continue
            if not os.path.exists(abs_path):
                continue
            seen.add(abs_path)
            resolved.append(abs_path)
        return resolved

    def _write_merged_xml(
        self, selected_candidates: Sequence[SuiteCandidate], output_file: str
    ) -> None:
        if not selected_candidates:
            raise ValueError("没有可合并的数据。")

        robot_attrs = self._read_robot_attrs(selected_candidates[0].file_path)
        merged_root = ET.Element("robot", attrib=robot_attrs)

        parent_cache: Dict[tuple[str, ...], ET.Element] = {}
        candidates_by_file: Dict[str, list[SuiteCandidate]] = {}
        file_order: list[str] = []
        for candidate in selected_candidates:
            file_candidates = candidates_by_file.get(candidate.file_path)
            if file_candidates is None:
                file_candidates = []
                candidates_by_file[candidate.file_path] = file_candidates
                file_order.append(candidate.file_path)
            file_candidates.append(candidate)

        for file_path in file_order:
            source_tree = parse_robot_tree(file_path)
            try:
                for candidate in candidates_by_file[file_path]:
                    suite_element = self._find_suite_element(source_tree, candidate)
                    if suite_element is None:
                        raise ValueError(
                            f"未能在 {candidate.file_path} 中定位 suite: {candidate.suite_path}"
                        )

                    parent = self._ensure_parent_path(
                        merged_root, candidate.suite_path_parts[:-1], parent_cache
                    )
                    parent.append(copy.deepcopy(suite_element))
            finally:
                root = source_tree.getroot()
                root.clear()
                del source_tree

        self._reindex_suite_ids(merged_root)
        merged_root.append(ET.Element("statistics"))
        merged_root.append(ET.Element("errors"))
        ET.ElementTree(merged_root).write(
            output_file, encoding="UTF-8", xml_declaration=True
        )

    @staticmethod
    def _read_robot_attrs(file_path: str) -> Dict[str, str]:
        generated = datetime.now().strftime("%Y%m%d %H:%M:%S.%f")[:-3]
        defaults = {
            "generator": "Robotframework-XMLMergeTool",
            "generated": generated,
            "rpa": "false",
            "schemaversion": "4",
        }
        tree = parse_robot_tree(file_path)
        root = tree.getroot()
        for key in ("generator", "rpa", "schemaversion"):
            if root.get(key):
                defaults[key] = root.get(key)  # type: ignore[assignment]
        defaults["generated"] = generated
        return defaults

    def _find_suite_element(
        self, source_tree: ET.ElementTree, candidate: SuiteCandidate
    ) -> ET.Element | None:
        root = source_tree.getroot()
        if candidate.suite_id:
            matches = root.xpath(".//suite[@id=$suite_id]", suite_id=candidate.suite_id)
            if matches:
                return matches[0]

        candidates = [node for node in root if node.tag == "suite"]
        for part in candidate.suite_path_parts:
            next_level: list[ET.Element] = []
            for node in candidates:
                if node.get("name") == part:
                    next_level.append(node)
            if next_level:
                candidates = next_level
                continue

            descendants: list[ET.Element] = []
            for node in candidates:
                descendants.extend(
                    [
                        child
                        for child in node
                        if child.tag == "suite" and child.get("name") == part
                    ]
                )
            candidates = descendants
            if not candidates:
                return None

        return candidates[0] if candidates else None

    @staticmethod
    def _ensure_parent_path(
        merged_root: ET.Element,
        parent_parts: Sequence[str],
        cache: Dict[tuple[str, ...], ET.Element],
    ) -> ET.Element:
        if not parent_parts:
            return merged_root

        current = merged_root
        path_prefix: list[str] = []
        for part in parent_parts:
            path_prefix.append(part)
            key = tuple(path_prefix)
            cached = cache.get(key)
            if cached is not None:
                current = cached
                continue

            existing = None
            for child in current:
                if child.tag == "suite" and child.get("name") == part:
                    existing = child
                    break

            if existing is None:
                existing = ET.Element("suite")
                existing.set("name", part)
                current.append(existing)

            cache[key] = existing
            current = existing

        return current

    def _write_conflict_report(
        self,
        all_candidates: Sequence[SuiteCandidate],
        selected_ids: set[str],
        suggestion: SelectionSuggestion,
        report_path: str,
    ) -> None:
        grouped: Dict[str, list[SuiteCandidate]] = defaultdict(list)
        for candidate in all_candidates:
            grouped[candidate.suite_path].append(candidate)

        default_selected = set(suggestion.selected_ids)
        os.makedirs(os.path.dirname(report_path), exist_ok=True)

        with open(report_path, "w", newline="", encoding="utf-8") as csv_file:
            writer = csv.DictWriter(
                csv_file,
                fieldnames=[
                    "suite_key",
                    "candidate_id",
                    "file_path",
                    "status",
                    "suite_endtime",
                    "test_endtime",
                    "effective_time",
                    "effective_time_source",
                    "default_selected",
                    "selected",
                    "decision",
                    "group_size",
                ],
            )
            writer.writeheader()

            for suite_key in sorted(grouped):
                group = grouped[suite_key]
                group_size = len(group)
                for candidate in group:
                    effective_time, time_source = self._policy.resolve_effective_time(
                        candidate
                    )
                    is_default_selected = candidate.candidate_id in default_selected
                    is_selected = candidate.candidate_id in selected_ids
                    decision = self._decision_reason(
                        is_selected=is_selected,
                        is_default_selected=is_default_selected,
                    )
                    writer.writerow(
                        {
                            "suite_key": suite_key,
                            "candidate_id": candidate.candidate_id,
                            "file_path": candidate.file_path,
                            "status": candidate.status,
                            "suite_endtime": candidate.suite_endtime or "",
                            "test_endtime": candidate.test_endtime or "",
                            "effective_time": effective_time.strftime(
                                "%Y-%m-%d %H:%M:%S.%f"
                            ),
                            "effective_time_source": time_source,
                            "default_selected": "yes" if is_default_selected else "no",
                            "selected": "yes" if is_selected else "no",
                            "decision": decision,
                            "group_size": str(group_size),
                        }
                    )

    @staticmethod
    def _decision_reason(*, is_selected: bool, is_default_selected: bool) -> str:
        if is_selected and is_default_selected:
            return "default-selected"
        if is_selected and not is_default_selected:
            return "user-selected"
        if not is_selected and is_default_selected:
            return "user-unselected-default"
        return "filtered-out"

    @staticmethod
    def _reindex_suite_ids(robot_root: ET.Element) -> None:
        top_suites = [child for child in robot_root if child.tag == "suite"]
        for top_index, suite in enumerate(top_suites, start=1):
            suite_id = f"s{top_index}"
            suite.set("id", suite_id)
            XmlMergeService._assign_suite_child_ids(suite, suite_id)

    @staticmethod
    def _assign_suite_child_ids(suite: ET.Element, suite_id: str) -> None:
        test_index = 0
        suite_index = 0
        for child in suite:
            if child.tag == "test":
                test_index += 1
                child.set("id", f"{suite_id}-t{test_index}")
            elif child.tag == "suite":
                suite_index += 1
                child_id = f"{suite_id}-s{suite_index}"
                child.set("id", child_id)
                XmlMergeService._assign_suite_child_ids(child, child_id)

    @staticmethod
    def _generate_robot_report(
        output_xml: str,
        log_path: str,
        report_path: str,
        output_dir: str,
        rebot_options: RebotOptions,
    ) -> None:
        reb = Rebot()
        execute_options = XmlMergeService._build_rebot_execute_options(
            output_dir=output_dir,
            log_path=log_path,
            report_path=report_path,
            rebot_options=rebot_options,
        )
        try:
            rc = reb.execute(output_xml, **execute_options)
        except Exception as exc:
            raise ValueError(f"Robot rebot 执行异常: {exc}") from exc
        if rc >= 252:
            raise ValueError(f"Robot rebot 执行失败，返回码: {rc}")
        if not os.path.exists(log_path) or not os.path.exists(report_path):
            raise ValueError("报告生成失败，未找到 log/report 文件。")

    @staticmethod
    def _generate_robot_log(
        output_xml: str,
        log_path: str,
        output_dir: str,
        rebot_options: RebotOptions,
    ) -> None:
        reb = Rebot()
        execute_options = XmlMergeService._build_rebot_execute_options(
            output_dir=output_dir,
            log_path=log_path,
            report_path=None,
            rebot_options=rebot_options,
        )
        try:
            rc = reb.execute(output_xml, **execute_options)
        except Exception as exc:
            raise ValueError(f"Robot rebot 执行异常: {exc}") from exc
        if rc >= 252:
            raise ValueError(f"Robot rebot 执行失败，返回码: {rc}")
        if not os.path.exists(log_path):
            raise ValueError("日志生成失败，未找到 log 文件。")

    @staticmethod
    def _resolve_rebot_streams() -> tuple[object, object]:
        stdout = sys.__stdout__ if sys.__stdout__ is not None else io.StringIO()
        stderr = sys.__stderr__ if sys.__stderr__ is not None else io.StringIO()
        return stdout, stderr

    @staticmethod
    def _build_rebot_execute_options(
        *,
        output_dir: str,
        log_path: str,
        report_path: str | None,
        rebot_options: RebotOptions,
    ) -> dict[str, object]:
        stdout, stderr = XmlMergeService._resolve_rebot_streams()
        options: dict[str, object] = {
            "log": log_path,
            "report": report_path if report_path else "NONE",
            "outputdir": output_dir,
            "stdout": stdout,
            "stderr": stderr,
        }
        options.update(rebot_options.to_rebot_kwargs(include_report=bool(report_path)))
        return options
