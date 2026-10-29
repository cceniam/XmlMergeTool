"""Native business layer differentials. Python is only a test oracle.

Usage: .venv/bin/python cpp/tests/test_app_compat.py build/desktop/xml-merge-driver
"""
from __future__ import annotations
from contextlib import contextmanager
import csv
import dataclasses
from datetime import datetime
import errno
import gc
import hashlib
import json
import os
import re
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from lxml import etree
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from xml_merge_tool.application.merge_service import XmlMergeService
from xml_merge_tool.application.models import RebotOptions
from xml_merge_tool.infrastructure.rf_reader import RFResultReader

DRIVER = Path(os.environ.get("XML_MERGE_DRIVER", ROOT / "build/desktop/xml-merge-driver")).resolve()
if __name__ == "__main__" and len(sys.argv) > 1:
    DRIVER = Path(sys.argv.pop(1)).resolve()


def xml(status="PASS", end="20260102 00:00:03.123456", root="根", source="C:\\路径\\用例.robot", child=True):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<robot generator="Robot 6.1.1" generated="20260101 10:00:00.000" rpa="false" schemaversion="4">
<suite name="{root}" id="s1" source="{source}">
<doc>保留文档</doc><meta name="版本">v1</meta>
<test name="=SUM(A1:A2) 测试" id="s1-t1"><kw name="Log" library="BuiltIn"><arg>信息</arg><msg level="INFO" timestamp="20260101 23:59:59.250">内容</msg><status status="PASS" starttime="20260101 23:59:59.250" endtime="20260101 23:59:59.251"/></kw><tag>标签</tag><status status="{status}" starttime="20260101 23:59:59.250" endtime="{end}">失败 &amp; 原因</status></test>
<test name="跳过" id="s1-t2"><status status="SKIP" starttime="20260102 00:00:04.000" endtime="20260102 00:00:04.000">跳过原因</status></test>
{('<suite name="子套件" id="s1-s1"><test name="子测试" id="s1-s1-t1"><status status="FAIL" starttime="20260102 00:00:00" endtime="20260102 00:00:01">子失败</status></test><status status="FAIL" starttime="20260102 00:00:00" endtime="20260102 00:00:01">套件失败</status></suite>') if child else ''}
<status status="{status}" starttime="20260101 23:59:58" endtime="{end}"/>
</suite><statistics/><errors/></robot>'''


def plain(value):
    return json.loads(json.dumps(value, ensure_ascii=False))


def semantic_xml(path):
    root = etree.parse(str(path)).getroot()
    root.attrib.pop("generated", None)
    def convert(n):
        return n.tag, dict(n.attrib), n.text or "", n.tail or "", [convert(c) for c in n]
    return convert(root)


@contextmanager
def deny_windows_directory_file_creation(directory):
    """Deny FILE_ADD_FILE on one test-owned directory, then restore its DACL."""
    import ctypes
    from ctypes import wintypes

    advapi32 = ctypes.WinDLL("advapi32", use_last_error=True)
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    pointer = ctypes.c_void_p
    pointer_pointer = ctypes.POINTER(pointer)
    advapi32.GetNamedSecurityInfoW.argtypes = (
        wintypes.LPWSTR, wintypes.DWORD, wintypes.DWORD,
        pointer_pointer, pointer_pointer, pointer_pointer, pointer_pointer, pointer_pointer,
    )
    advapi32.GetNamedSecurityInfoW.restype = wintypes.DWORD
    advapi32.SetNamedSecurityInfoW.argtypes = (
        wintypes.LPWSTR, wintypes.DWORD, wintypes.DWORD, pointer, pointer, pointer, pointer,
    )
    advapi32.SetNamedSecurityInfoW.restype = wintypes.DWORD
    advapi32.GetSecurityDescriptorControl.argtypes = (
        pointer, ctypes.POINTER(wintypes.WORD), ctypes.POINTER(wintypes.DWORD),
    )
    advapi32.GetSecurityDescriptorControl.restype = wintypes.BOOL
    advapi32.ConvertStringSecurityDescriptorToSecurityDescriptorW.argtypes = (
        wintypes.LPCWSTR, wintypes.DWORD, pointer_pointer, ctypes.POINTER(wintypes.DWORD),
    )
    advapi32.ConvertStringSecurityDescriptorToSecurityDescriptorW.restype = wintypes.BOOL
    advapi32.GetSecurityDescriptorDacl.argtypes = (
        pointer, ctypes.POINTER(wintypes.BOOL), pointer_pointer, ctypes.POINTER(wintypes.BOOL),
    )
    advapi32.GetSecurityDescriptorDacl.restype = wintypes.BOOL
    kernel32.LocalFree.argtypes = (pointer,)
    kernel32.LocalFree.restype = pointer

    def check_error(error):
        if error:
            raise ctypes.WinError(error)

    def check_bool(success):
        if not success:
            raise ctypes.WinError(ctypes.get_last_error())

    original_sd, original_dacl = pointer(), pointer()
    restricted_sd, restricted_dacl = pointer(), pointer()
    control, revision = wintypes.WORD(), wintypes.DWORD()
    present, defaulted = wintypes.BOOL(), wintypes.BOOL()
    applied = False
    try:
        check_error(advapi32.GetNamedSecurityInfoW(
            str(directory), 1, 4, None, None, ctypes.byref(original_dacl), None,
            ctypes.byref(original_sd)))
        check_bool(advapi32.GetSecurityDescriptorControl(
            original_sd, ctypes.byref(control), ctypes.byref(revision)))
        # The directory is newly created for this test. Deny everyone file
        # creation even for administrator runners; retain traversal and the
        # WRITE_DAC permission needed to restore the original descriptor.
        check_bool(advapi32.ConvertStringSecurityDescriptorToSecurityDescriptorW(
            "D:P(D;;0x2;;;WD)(A;;FA;;;WD)", 1, ctypes.byref(restricted_sd), None))
        check_bool(advapi32.GetSecurityDescriptorDacl(
            restricted_sd, ctypes.byref(present), ctypes.byref(restricted_dacl),
            ctypes.byref(defaulted)))
        check_error(advapi32.SetNamedSecurityInfoW(
            str(directory), 1, 4 | 0x80000000, None, None, restricted_dacl, None))
        applied = True
        yield
    finally:
        try:
            if applied:
                # Restore both the ACL and its original inheritance protection.
                inheritance = 0x80000000 if control.value & 0x1000 else 0x20000000
                check_error(advapi32.SetNamedSecurityInfoW(
                    str(directory), 1, 4 | inheritance, None, None, original_dacl, None))
        finally:
            if restricted_sd:
                kernel32.LocalFree(restricted_sd)
            if original_sd:
                kernel32.LocalFree(original_sd)


class NativeBusinessTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="native-业务-")
        self.dir = Path(self.tmp.name)
        self.service = XmlMergeService()

    def tearDown(self):
        # The unchanged Python oracle retains exception/iterparse cycles for
        # malformed roots. Release their open input handles before Windows
        # removes the fixtures; native subprocesses have already exited.
        gc.collect()
        self.tmp.cleanup()

    def write(self, name, text):
        path = self.dir / name
        path.write_text(text, encoding="utf-8")
        os.utime(path, (1767225600.25, 1767225600.25))
        return path

    def native(self, command, *args, fail=False, env=None, cwd=None):
        result = subprocess.run([str(DRIVER), command, *map(str, args)], capture_output=True, text=True, encoding="utf-8", timeout=120, env=env, cwd=cwd)
        if fail:
            self.assertNotEqual(result.returncode, 0)
            return result.stderr
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def compare_load(self, paths, *, compare_selection=True):
        native = self.native("load", *paths)
        result = self.service.load_candidates(list(map(str, paths)))
        self.assertEqual(native["candidates"], plain([dataclasses.asdict(c) for c in result.candidates]))
        self.assertEqual(native["parsed_file_count"], result.parsed_file_count)
        self.assertEqual(native["failed_file_count"], result.failed_file_count)
        self.assertEqual(native["successful_file_paths"], list(result.successful_file_paths))
        if compare_selection:
            suggestion = self.service.suggest_selection(result.candidates)
            self.assertEqual(native["selected_ids"], list(suggestion.selected_ids))
            self.assertEqual(native["conflicts"], plain([dataclasses.asdict(c) for c in suggestion.conflicts]))
        return native, result

    def test_streaming_fields_preorder_direct_tests_inherited_source(self):
        a = self.write("输入.xml", xml("FAIL"))
        self.compare_load([a])

    def test_parallel_order_normalization_and_failure_isolation(self):
        files = [self.write(f"{i}.xml", xml(root=f"Root{i}")) for i in range(8)]
        files += [self.write("bad.xml", "not xml"), self.write("wrong.xml", "<other/>")]
        files += [files[0], self.write("ignored.txt", xml()), self.dir / "missing.xml"]
        native, _ = self.compare_load(files)
        self.assertEqual(native["failed_file_count"], 2)
        serial = self.native("load", "--workers", 1, *files)
        parallel = self.native("load", "--workers", 3, *files)
        self.assertEqual(serial, parallel)

    def test_recover_sanitize_and_no_ids(self):
        inputs = {
            "trailing.xml": xml() + "junk trailing",
            "prefix.xml": "garbage\0\0" + xml() + "\0footer",
            "truncated.xml": xml().replace("</robot>", ""),
            "unnamed.xml": '<robot><suite><test><status status="fail" endtime="bad"/></test></suite></robot>',
            "no-id.xml": xml().replace(' id="s1"', '').replace(' id="s1-s1"', ''),
            "empty.xml": '<robot><suite name="empty"><status status="PASS"/></suite></robot>',
        }
        self.compare_load([self.write(name, value) for name, value in inputs.items()])

    def test_namespaced_attributes_do_not_change_selected_suite(self):
        for index, id_attrs in enumerate((
            'id="s1-s1" meta:id="s1-s2"',
            'meta:id="s1-s2" id="s1-s1"',
            'meta:id="s1-s2"',
        )):
            for recover in (False, True):
                with self.subTest(id_attrs=id_attrs, recover=recover):
                    name_attrs = 'meta:name="B" name="A"' if index == 1 else 'name="A" meta:name="B"'
                    test_name_attrs = 'meta:name="B test" name="A test"' if index == 1 else 'name="A test" meta:name="B test"'
                    status_attrs = 'meta:status="PASS" status="FAIL"' if index == 1 else 'status="FAIL" meta:status="PASS"'
                    content = f'''<robot generator="Robot 6.1.1" schemaversion="4" xmlns:meta="urn:metadata">
<suite name="Root" id="s1"><suite {name_attrs} {id_attrs}>
<test {test_name_attrs}><status {status_attrs}>Important failure</status></test>
<status {status_attrs}/></suite>
<suite name="B" id="s1-s2"><test name="B test"><status status="PASS"/></test><status status="PASS"/></suite>
<status status="FAIL"/></suite><statistics/><errors/></robot>'''
                    # A trailing token makes the native streaming parse fall
                    # back to the recovery tree scanner.
                    if recover:
                        content += "trailing junk"
                    source = self.write(f"namespaced-{index}-{recover}.xml", content)
                    native, result = self.compare_load([source])
                    self.assertEqual(native["candidates"][0]["suite_name"], "A")
                    self.assertEqual(native["candidates"][0]["status"], "FAIL")
                    candidate = result.candidates[0]
                    expected = self.dir / "expected.xml"
                    self.service.export_single_candidate_xml(candidate, str(expected))
                    exported = self.native("xml", self.dir / "selected.xml", "--select", "0", source)
                    self.assertEqual(semantic_xml(exported["output_path"]), semantic_xml(expected))
                    merged = self.native("merge", self.dir / f"merged-{index}-{recover}", "--select", "0", source)
                    expected_merge = self.dir / "expected-merge.xml"
                    self.service._write_merged_xml([candidate], str(expected_merge))
                    self.assertEqual(semantic_xml(merged["merged_xml_path"]), semantic_xml(expected_merge))
                    tree = etree.parse(merged["merged_xml_path"])
                    self.assertEqual(tree.xpath("//test/@name"), ["A test"])
                    self.assertEqual(tree.xpath("//test/status/@status"), ["FAIL"])
                    model = json.loads(subprocess.check_output(
                        ["node", str(ROOT / "cpp/tests/extract_model.js"), merged["report_path"]],
                        text=True, encoding="utf-8", timeout=120))
                    total = model["output"]["stats"][0][0]
                    self.assertEqual((total["pass"], total["fail"]), (0, 1))

    def test_unicode_whitespace_candidates_conflicts_and_merge(self):
        for marker in ("\u00a0", "\u3000"):
            with self.subTest(codepoint=hex(ord(marker))):
                paths = [
                    self.write(f"{ord(marker)}-old.xml", xml("FAIL", root="Suite", child=False)),
                    self.write(f"{ord(marker)}-rerun.xml", xml("PASS", root=marker + "Suite" + marker, child=False)
                               .replace('name="跳过"', f'name="{marker}跳过{marker}"')
                               .replace('source="C:', f'source="{marker}C:')
                               .replace('用例.robot"', f'用例.robot{marker}"')),
                ]
                native, result = self.compare_load(paths)
                self.assertEqual(len(native["selected_ids"]), 1)
                self.assertEqual(len(native["conflicts"]), 1)
                self.assertEqual(native["candidates"][1]["suite_name"], "Suite")
                self.assertEqual(native["candidates"][1]["test_names"][1], "跳过")
                self.assertTrue(native["selected_ids"][0].startswith(str(paths[1])))
                selected = [c for c in result.candidates if c.candidate_id in native["selected_ids"]]
                expected_export = self.dir / "unicode-export-expected.xml"
                self.service.export_single_candidate_xml(selected[0], str(expected_export))
                exported = self.native("xml", self.dir / "unicode-export.xml", *paths)
                self.assertEqual(semantic_xml(exported["output_path"]), semantic_xml(expected_export))
                expected = self.dir / "unicode-expected.xml"
                self.service._write_merged_xml(selected, str(expected))
                merged = self.native("merge", self.dir / f"unicode-{ord(marker)}", *paths)
                self.assertEqual(merged["selected_count"], 1)
                self.assertEqual(semantic_xml(merged["merged_xml_path"]), semantic_xml(expected))
                self.assertEqual(len(etree.parse(merged["merged_xml_path"]).xpath("//test")), 2)

    def test_unicode_whitespace_status_times_and_messages(self):
        for marker in ("\u00a0", "\u3000"):
            with self.subTest(codepoint=hex(ord(marker))):
                content = xml("FAIL", child=False).replace('status="FAIL"', f'status="{marker}FAIL{marker}"')
                content = content.replace('="2026', f'="{marker}2026').replace('03.123456"', f'03.123456{marker}"')
                content = content.replace("失败 &amp; 原因", marker + "失败 &amp; 原因" + marker)
                self.compare_load([self.write(f"unicode-status-{ord(marker)}.xml", content)])

    def test_selection_priority_time_fallback_tie_and_start_is_display_only(self):
        files = [
            self.write("01-fail.xml", xml("FAIL", "20290101 00:00:00")),
            self.write("02-pass.xml", xml("PASS", "20260102 00:00:03")),
            self.write("03-skip.xml", xml("SKIP", "20260102 00:00:04")),
            self.write("04-skip-tie.xml", xml("SKIP", "20260102 00:00:04")),
            self.write("05-bad-suite-time.xml", xml("PASS").replace('endtime="20260102 00:00:03.123456"/>', 'endtime="invalid"/>')),
        ]
        self.compare_load(files)
        t = '<robot><suite name="fallback"><test name="one"><status status="PASS"/></test><status status="PASS" starttime="20990101 00:00:00"/></suite></robot>'
        a, b = self.write("a.xml", t), self.write("b.xml", t.replace("20990101", "20200101"))
        os.utime(b, (1767225601.25, 1767225601.25))
        native, _ = self.compare_load([a, b])
        self.assertTrue(native["selected_ids"][0].startswith(str(b)))

    def test_file_mtime_preserves_subsecond_precision(self):
        text = '<robot><suite name="mtime"><test name="one"><status status="PASS"/></test><status status="PASS"/></suite></robot>'
        path = self.write("fractional.xml", text)
        nanoseconds = 1767225600987654300
        os.utime(path, ns=(nanoseconds, nanoseconds))
        native, _ = self.compare_load([path])
        self.assertGreater(native["candidates"][0]["file_mtime"], 1767225600.9)

    def test_negative_file_mtime_candidates_and_epoch_fallback(self):
        text = '<robot><suite name="mtime"><test name="one"><status status="PASS"/></test><status status="PASS"/></suite></robot>'
        path = self.write("before-epoch.xml", text)
        nanoseconds = -85412345700
        os.utime(path, ns=(nanoseconds, nanoseconds))
        native, reference = self.compare_load([path], compare_selection=False)
        self.assertEqual(native["candidates"][0]["file_mtime"], path.stat().st_mtime)
        self.assertLess(native["candidates"][0]["file_mtime"], 0)
        self.assertEqual(native["selected_ids"], [reference.candidates[0].candidate_id])
        effective, source = self.service._policy.resolve_effective_time(reference.candidates[0])
        self.assertEqual((effective, source), (datetime.fromtimestamp(0), "file_mtime"))
        try:
            effective.timestamp()
        except OSError as error:
            # Some Windows CRT/Python combinations cannot round-trip the epoch
            # via naive datetime.timestamp(). Candidate parsing and native
            # selection remain checked above; record the oracle's limitation.
            self.assertEqual((os.name, error.errno), ("nt", 22))
            with self.assertRaises(OSError) as failed:
                self.service.suggest_selection(reference.candidates)
            self.assertEqual(failed.exception.errno, 22)
        else:
            suggestion = self.service.suggest_selection(reference.candidates)
            self.assertEqual(native["selected_ids"], list(suggestion.selected_ids))

    def test_mtime_fallback_rounds_to_python_microseconds_for_ties_and_csv(self):
        text = '<robot><suite name="mtime"><test name="one"><status status="PASS"/></test><status status="PASS"/></suite></robot>'
        paths = [self.write("a.xml", text), self.write("z.xml", text)]
        for path, nanoseconds in zip(paths, (1767225600123456100, 1767225600123456400)):
            os.utime(path, ns=(nanoseconds, nanoseconds))
        native, result = self.compare_load(paths)
        self.assertTrue(native["selected_ids"][0].startswith(str(paths[0])))
        out = self.native("merge", self.dir / "native", *paths)
        reference = self.dir / "expected.csv"
        suggestion = self.service.suggest_selection(result.candidates)
        self.service._write_conflict_report(result.candidates, set(suggestion.selected_ids),
                                            suggestion, str(reference))
        with open(out["conflict_report_path"], newline="", encoding="utf-8") as actual, reference.open(newline="", encoding="utf-8") as expected:
            self.assertEqual(list(csv.DictReader(actual)), list(csv.DictReader(expected)))

    def test_microsecond_duration_cross_day_and_invalid_calendar(self):
        times = [("20260101 23:59:59.123456", "20260102 00:00:00.123457"), ("20260101 00:00:01", "20260101 00:00:01"), ("20260230 00:00:00", "20260301 00:00:00"), ("20260102 00:00:00", "20260101 00:00:00"), ("", "")]
        tests = ''.join(f'<test name="{i}"><status status="UNKNOWN" starttime="{start}" endtime="{end}"/></test>' for i, (start,end) in enumerate(times))
        self.compare_load([self.write("durations.xml", f'<robot><suite name="Root">{tests}</suite></robot>')])

    def test_full_subtrees_parent_child_duplication_and_conflict_csv(self):
        paths = [self.write("1.xml", xml("FAIL")), self.write("2.xml", xml("PASS"))]
        result = self.service.load_candidates(list(map(str, paths)))
        indices = [2, 0, 1, 2]
        selected = [result.candidates[i] for i in indices]
        out = self.native("merge", self.dir / "native", "--select", ','.join(map(str, indices)), *paths)
        reference = self.dir / "reference.xml"
        self.service._write_merged_xml(selected, str(reference))
        self.assertEqual(semantic_xml(out["merged_xml_path"]), semantic_xml(reference))
        expected_csv = self.dir / "reference.csv"
        self.service._write_conflict_report(result.candidates, {c.candidate_id for c in selected}, self.service.suggest_selection(result.candidates), str(expected_csv))
        with open(out["conflict_report_path"], newline="", encoding="utf-8") as actual, expected_csv.open(newline="", encoding="utf-8") as expected:
            self.assertEqual(list(csv.DictReader(actual)), list(csv.DictReader(expected)))
        for key in ("log_path", "report_path"):
            self.assertGreater(Path(out[key]).stat().st_size, 10000)
        self.assertFalse(list((self.dir / "native").glob(".merge-*")))

    def test_schema3_and_source_unchanged(self):
        source = ROOT / "cpp/tests/fixtures/schema3.xml"
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        self.compare_load([source])
        result = self.native("xml", self.dir / "single.xml", source)
        native_tree = etree.parse(result["output_path"])
        self.assertEqual(native_tree.getroot().get("schemaversion"), "3")
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
        self.native("xml", source, source, fail=True)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_single_log_removes_temp_xml_and_xml_matches_oracle(self):
        source = self.write("单独.xml", xml())
        result = self.service.load_candidates([str(source)])
        out = self.native("xml", self.dir / "导出.xml", "--select", "1", source)
        reference = self.dir / "参考.xml"
        self.service.export_single_candidate_xml(result.candidates[1], str(reference))
        self.assertEqual(semantic_xml(out["output_path"]), semantic_xml(reference))
        log = self.native("log", self.dir / "日志", "--select", "1", source)
        self.assertTrue(Path(log["output_path"]).is_file())
        self.assertEqual([p.name for p in (self.dir / "日志").iterdir()], ["single-log.html"])

    def test_single_log_omits_report_only_settings_like_python(self):
        content = re.sub(
            r"(\d{8} \d{2}:\d{2}:\d{2})(?:\.(\d+))?",
            lambda m: m[1] + "." + (m[2] or "").ljust(3, "0")[:3], xml())
        source = self.write("single-options.xml", content)
        candidate = self.service.load_candidates([str(source)]).candidates[0]
        reference = self.service.generate_single_candidate_log(
            candidate, str(self.dir / "reference"),
            RebotOptions(suite_stat_level=1, report_background="invalid-color-count:1:2:3"),
        )
        native = self.native("log", self.dir / "native", "--suitestatlevel", "1",
                             "--reportbackground", "invalid-color-count:1:2:3", source)["output_path"]
        values = []
        for report_path in (reference, native):
            data = json.loads(subprocess.check_output(
                ["node", str(ROOT / "cpp/tests/extract_model.js"), report_path],
                text=True, encoding="utf-8", timeout=120))
            values.append(data["output"]["stats"])
        self.assertEqual(values[0], values[1])
        self.assertEqual(len(values[1][2]), 2)

    def test_excel_contents_order_format_and_formula_safety(self):
        paths = [self.write("1.xml", xml("FAIL")), self.write("2.xml", xml("PASS", source="/var/tests/second.robot"))]
        result = self.service.load_candidates(list(map(str, paths)))
        indices = [3, 0, 3]
        ids = [result.candidates[i].candidate_id for i in indices]
        expected = self.dir / "reference.xlsx"
        self.service.export_selected_tests_excel(result.candidates, ids, str(expected))
        out = self.native("excel", self.dir / "实际.xlsx", "--select", "3,0,3", *paths)
        self.assertEqual(out["selected_suite_count"], 2)
        for workbook_path in (expected, Path(out["output_path"])):
            book = load_workbook(workbook_path)
            sheet = book.active
            values = [[(c.value, c.data_type if c.value is not None else None, c.number_format) for c in row] for row in sheet]
            layout = (book.sheetnames, str(sheet.freeze_panes), sheet.auto_filter.ref, sheet.sheet_view.showGridLines, sheet.sheet_view.zoomScale)
            if workbook_path == expected:
                expected_values, expected_layout = values, layout
            else:
                self.assertEqual(values, expected_values)
                self.assertEqual(layout, expected_layout)
                self.assertEqual(len(sheet.conditional_formatting), 1)
            book.close()
        self.assertFalse(list(self.dir.glob(".实际-*")))

    def test_excel_long_text_invalid_duration_and_atomic_failure(self):
        source = self.write("long.xml", xml("FAIL", child=False).replace("失败 &amp; 原因", "长" * 33000).replace('endtime="20260102 00:00:03.123456"', 'endtime="bad"'))
        out = self.native("excel", self.dir / "长.xlsx", source)
        book = load_workbook(out["output_path"])
        self.assertEqual(len(book.active["E2"].value), 32767)
        self.assertTrue(book.active["E2"].value.endswith("\n[内容过长，已截断]"))
        self.assertIsNone(book.active["D2"].value)
        book.close()
        existing = self.dir / "keep.xlsx"
        existing.write_bytes(b"keep")
        self.native("excel", existing, "--select", "", source, fail=True)
        self.assertEqual(existing.read_bytes(), b"keep")
        self.assertFalse(list(self.dir.glob(".keep-*")))

    @unittest.skipUnless(os.name == "nt", "Windows UTF-16 environment and directory permissions")
    def test_excel_unicode_temp_with_unwritable_working_directory(self):
        protected_cwd = self.dir / "unwritable-cwd"
        protected_cwd.mkdir()
        source = self.write("临时文件输入.xml", xml(child=False))
        source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
        temp_dir = self.dir / "中文临时目录🧪"
        temp_dir.mkdir()
        # This fixture must reproduce the ANSI helper's loss of the temp path,
        # rather than silently passing on a UTF-8 ACP Windows configuration.
        ansi_temp = str(temp_dir).encode("mbcs", errors="replace").decode("mbcs")
        self.assertNotEqual(ansi_temp, str(temp_dir), "Regression requires a non-UTF-8 ANSI code page")
        environment = dict(os.environ, TMP=str(temp_dir), TEMP=str(temp_dir))
        output = self.dir / "中文结果.xlsx"
        with deny_windows_directory_file_creation(protected_cwd):
            probe = protected_cwd / "creation-must-fail.tmp"
            with self.assertRaises(PermissionError) as denied:
                with probe.open("xb") as stream:
                    stream.write(b"denied")
            # Python's CRT file open reports EACCES without a Win32 error code.
            self.assertEqual(denied.exception.errno, errno.EACCES)
            self.assertIn(getattr(denied.exception, "winerror", None), (None, 5))
            self.assertFalse(probe.exists())
            result = self.native("excel", output, source, cwd=protected_cwd, env=environment)
        # Confirm the fixture restored its original ACL before normal cleanup.
        restored = protected_cwd / "creation-restored.tmp"
        restored.write_bytes(b"restored")
        restored.unlink()
        self.assertEqual(Path(result["output_path"]), output)
        book = load_workbook(output)
        try:
            sheet = book.active
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A2"].value, "=SUM(A1:A2) 测试")
            self.assertEqual(sheet["A2"].data_type, "s")
            self.assertEqual(sheet["C2"].value, "PASS")
            self.assertEqual(sheet["D2"].number_format, "[h]:mm:ss.000")
        finally:
            book.close()
        self.assertEqual(list(temp_dir.iterdir()), [])
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_sha)

    def test_source_links_are_relative_to_final_report_directory(self):
        import base64
        import zlib
        robot_file = self.dir / "actual.robot"
        robot_file.write_text("*** Test Cases ***\nTest\n    No Operation\n")
        source = self.write("source.xml", xml(source=str(robot_file), child=False))
        merge = self.native("merge", self.dir / "output", source)
        log = self.native("log", self.dir / "single", source)
        for report_path in (merge["log_path"], log["output_path"]):
            data = json.loads(subprocess.check_output(["node", str(ROOT / "cpp/tests/extract_model.js"), report_path], text=True, encoding="utf-8", timeout=120))
            strings = [s[1:] if s.startswith("*") else zlib.decompress(base64.b64decode(s)).decode() for s in data["output"]["strings"]]
            relative = strings[data["output"]["suite"][2]]
            self.assertEqual((Path(report_path).parent / relative).resolve(), robot_file.resolve())

    def test_split_log_sidecar_source_protection_and_cleanup(self):
        source = self.write("original.xml", xml())
        output = self.dir / "output"
        output.mkdir()
        sidecar = output / "single-log-1.js"
        os.link(source, sidecar)
        before = source.read_bytes()
        self.native("log", output, "--splitlog", source, fail=True)
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((output / "single-log.html").exists())
        self.assertFalse(list(output.glob(".single-log-*")))
        sidecar.unlink()
        result = self.native("log", output, "--splitlog", source)
        self.assertTrue(Path(result["output_path"]).is_file())
        self.assertTrue(sidecar.is_file())

    def test_equivalent_source_path_protection(self):
        source = self.write("original.xml", xml())
        alias = self.dir / "alias.xml"
        alias.symlink_to(source)
        # The XML and its fallback timestamp must both come from the target,
        # while candidate identity retains the user-supplied alias path.
        self.compare_load([alias])
        before = source.read_bytes()
        self.native("xml", alias, source, fail=True)
        self.assertEqual(source.read_bytes(), before)
        xlsx_alias = self.dir / "source.xlsx"
        os.link(source, xlsx_alias)
        self.native("excel", xlsx_alias, source, fail=True)
        self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
