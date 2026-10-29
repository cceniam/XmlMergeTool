from __future__ import annotations

import shutil
import sys
import unittest
import uuid
from contextlib import contextmanager
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

SRC_ROOT = Path(__file__).resolve().parents[1] / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from xml_merge_tool.application.merge_service import XmlMergeService
from xml_merge_tool.application.models import SuiteCandidate, TestCaseResult as CaseResult


@contextmanager
def _workspace_temp_dir():
    root = Path(__file__).resolve().parents[1] / "tmp" / f"test-excel-{uuid.uuid4().hex}"
    root.mkdir(parents=True, exist_ok=True)
    try:
        yield root
    finally:
        shutil.rmtree(root, ignore_errors=True)


def _candidate(
    candidate_id: str,
    *,
    file_name: str,
    robot_source: str | None,
    tests: tuple[CaseResult, ...],
) -> SuiteCandidate:
    return SuiteCandidate(
        candidate_id=candidate_id,
        file_path=f"D:/logs/{file_name}",
        file_name=file_name,
        suite_id=f"suite-{candidate_id}",
        suite_name=f"Suite-{candidate_id}",
        suite_path=f"Root / Suite-{candidate_id}",
        suite_path_parts=("Root", f"Suite-{candidate_id}"),
        status="FAIL" if any(test.status == "FAIL" for test in tests) else "PASS",
        suite_endtime="20260101 10:00:01.000",
        test_endtime="20260101 10:00:01.000",
        file_mtime=0.0,
        robot_source=robot_source,
        test_results=tests,
    )


class TestExcelExport(unittest.TestCase):
    def setUp(self):
        self.service = XmlMergeService()

    def test_exports_selected_tests_in_candidate_and_xml_order(self):
        candidates = [
            _candidate(
                "candidate-a",
                file_name="a-output.xml",
                robot_source=r"D:\\robot\\登录测试.robot",
                tests=(
                    CaseResult(name="登录成功", status="PASS", duration_seconds=1.25),
                    CaseResult(
                        name="=HYPERLINK(\"https://example.invalid\")",
                        status="FAIL",
                        failure_reason="账号或密码错误",
                        duration_seconds=2.0,
                    ),
                ),
            ),
            _candidate(
                "candidate-b",
                file_name="b-output.xml",
                robot_source="/opt/robot/退出测试.robot",
                tests=(CaseResult(name="退出", status="SKIP"),),
            ),
            _candidate(
                "candidate-c",
                file_name="fallback-output.xml",
                robot_source=None,
                tests=(CaseResult(name="未勾选", status="PASS"),),
            ),
        ]

        with _workspace_temp_dir() as root:
            output_path = root / "selected-tests.xlsx"
            result = self.service.export_selected_tests_excel(
                candidates,
                ["candidate-b", "candidate-a"],
                str(output_path),
            )

            self.assertEqual(result.output_path, str(output_path.resolve()))
            self.assertEqual(result.selected_suite_count, 2)
            self.assertEqual(result.test_count, 3)

            workbook = load_workbook(output_path, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["测试结果"])
                worksheet = workbook["测试结果"]
                self.assertEqual(
                    list(worksheet.values),
                    [
                        ("测试名称", "Robot 文件名", "状态", "耗时", "失败原因"),
                        (
                            "登录成功",
                            "登录测试.robot",
                            "PASS",
                            timedelta(seconds=1.25),
                            None,
                        ),
                        (
                            '=HYPERLINK("https://example.invalid")',
                            "登录测试.robot",
                            "FAIL",
                            timedelta(seconds=2),
                            "账号或密码错误",
                        ),
                        ("退出", "退出测试.robot", "SKIP", None, None),
                    ],
                )
                self.assertEqual(worksheet["A3"].data_type, "s")
                self.assertNotIsInstance(worksheet["D2"].value, str)
                self.assertEqual(worksheet["D2"].number_format, "[h]:mm:ss.000")
                self.assertEqual(worksheet.freeze_panes, "A2")
                self.assertEqual(worksheet.auto_filter.ref, "A1:E4")
                self.assertFalse(worksheet.sheet_view.showGridLines)
                self.assertEqual(worksheet.column_dimensions["D"].width, 18)
                self.assertEqual(worksheet.column_dimensions["E"].width, 80)
                self.assertEqual(len(worksheet.conditional_formatting), 1)
                self.assertEqual(worksheet.sheet_view.zoomScale, 90)
                self.assertEqual(
                    worksheet.sheet_properties.tabColor.rgb[-6:], "1F4E78"
                )
                self.assertEqual(worksheet["A1"].fill.fgColor.rgb[-6:], "1F4E78")
                self.assertTrue(worksheet["A1"].font.bold)
                self.assertEqual(worksheet["A2"].fill.fgColor.rgb[-6:], "FFFFFF")
                self.assertEqual(worksheet["A3"].fill.fgColor.rgb[-6:], "EAF2F8")
                self.assertEqual(worksheet["A4"].fill.fgColor.rgb[-6:], "FFFFFF")
                self.assertEqual(
                    {
                        worksheet.cell(row=3, column=column).fill.fgColor.rgb[-6:]
                        for column in range(1, 6)
                    },
                    {"EAF2F8"},
                )
                self.assertEqual(worksheet["A2"].border.bottom.style, "thin")
            finally:
                workbook.close()

    def test_uses_input_xml_name_when_robot_source_is_missing(self):
        candidate = _candidate(
            "candidate-a",
            file_name="fallback-output.xml",
            robot_source="",
            tests=(CaseResult(name="case", status="PASS"),),
        )
        with _workspace_temp_dir() as root:
            output_path = root / "fallback.xlsx"
            self.service.export_selected_tests_excel(
                [candidate], [candidate.candidate_id], str(output_path)
            )
            workbook = load_workbook(output_path)
            try:
                self.assertEqual(workbook["测试结果"]["B2"].value, "fallback-output.xml")
            finally:
                workbook.close()

    def test_rejects_empty_invalid_and_testless_selection(self):
        candidate = _candidate(
            "candidate-a",
            file_name="a.xml",
            robot_source="a.robot",
            tests=tuple(),
        )
        with _workspace_temp_dir() as root:
            output_path = root / "result.xlsx"
            with self.assertRaisesRegex(ValueError, "至少勾选"):
                self.service.export_selected_tests_excel([candidate], [], str(output_path))
            with self.assertRaisesRegex(ValueError, "已失效"):
                self.service.export_selected_tests_excel(
                    [candidate], ["missing"], str(output_path)
                )
            with self.assertRaisesRegex(ValueError, "没有可导出"):
                self.service.export_selected_tests_excel(
                    [candidate], [candidate.candidate_id], str(output_path)
                )
            self.assertFalse(output_path.exists())

    def test_atomic_replace_failure_preserves_existing_file_and_cleans_temp(self):
        candidate = _candidate(
            "candidate-a",
            file_name="a.xml",
            robot_source="a.robot",
            tests=(CaseResult(name="case", status="PASS"),),
        )
        with _workspace_temp_dir() as root:
            output_path = root / "result.xlsx"
            output_path.write_bytes(b"existing-content")

            with patch(
                "xml_merge_tool.application.merge_service.os.replace",
                side_effect=PermissionError("locked"),
            ):
                with self.assertRaisesRegex(ValueError, "无法写入"):
                    self.service.export_selected_tests_excel(
                        [candidate], [candidate.candidate_id], str(output_path)
                    )

            self.assertEqual(output_path.read_bytes(), b"existing-content")
            self.assertEqual(list(root.glob(".*.xlsx")), [])

    def test_unwritable_output_directory_has_localized_error(self):
        candidate = _candidate(
            "candidate-a",
            file_name="a.xml",
            robot_source="a.robot",
            tests=(CaseResult(name="case", status="PASS"),),
        )
        with _workspace_temp_dir() as root:
            output_path = root / "blocked" / "result.xlsx"
            with patch(
                "xml_merge_tool.application.merge_service.os.makedirs",
                side_effect=PermissionError("denied"),
            ):
                with self.assertRaisesRegex(ValueError, "无法写入"):
                    self.service.export_selected_tests_excel(
                        [candidate], [candidate.candidate_id], str(output_path)
                    )
            self.assertFalse(output_path.exists())

    def test_truncates_overlong_text_with_visible_marker(self):
        value = self.service._excel_text("x" * 40_000)
        self.assertEqual(len(value), 32_767)
        self.assertTrue(value.endswith("[内容过长，已截断]"))

    def test_duration_cells_support_over_24_hours_zero_and_missing_values(self):
        candidate = _candidate(
            "candidate-a",
            file_name="a.xml",
            robot_source="a.robot",
            tests=(
                CaseResult(name="long", status="PASS", duration_seconds=90_061.234),
                CaseResult(name="zero", status="PASS", duration_seconds=0.0),
                CaseResult(name="missing", status="PASS", duration_seconds=None),
                CaseResult(name="negative", status="PASS", duration_seconds=-1.0),
            ),
        )
        with _workspace_temp_dir() as root:
            output_path = root / "durations.xlsx"
            self.service.export_selected_tests_excel(
                [candidate], [candidate.candidate_id], str(output_path)
            )

            workbook = load_workbook(output_path)
            try:
                worksheet = workbook["测试结果"]
                self.assertEqual(worksheet["D2"].value, timedelta(seconds=90_061.234))
                self.assertEqual(worksheet["D3"].value, timedelta(0))
                self.assertIsNone(worksheet["D4"].value)
                self.assertIsNone(worksheet["D5"].value)
                for row in range(2, 6):
                    self.assertEqual(
                        worksheet.cell(row=row, column=4).number_format,
                        "[h]:mm:ss.000",
                    )
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
