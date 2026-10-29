from __future__ import annotations

import os
import shutil
import sys
import unittest
import uuid
from contextlib import contextmanager
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

SRC_ROOT = Path(__file__).resolve().parents[1] / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt6.QtCore import QSettings, Qt
from PyQt6.QtWidgets import QApplication

from xml_merge_tool.ui.controller import MainController


@contextmanager
def _workspace_temp_dir():
    root = (
        Path(__file__).resolve().parents[1]
        / "tmp"
        / f"test-controller-{uuid.uuid4().hex}"
    )
    root.mkdir(parents=True, exist_ok=True)
    try:
        yield root
    finally:
        shutil.rmtree(root, ignore_errors=True)


def _write_robot_output(
    file_path: Path,
    *,
    root_name: str,
    leaf_name: str,
    status: str = "PASS",
    endtime: str = "20260101 10:00:01.000",
) -> None:
    starttime = "20260101 10:00:00.000"
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<robot generator="Robot 6.1.1" generated="20260101 10:00:00.000" rpa="false" schemaversion="4">
  <suite id="s1" name="{root_name}" source="{leaf_name}.robot">
    <suite id="s1-s1" name="{leaf_name}" source="{leaf_name}.robot">
      <test id="s1-s1-t1" name="case1" line="1">
        <status status="{status}" starttime="{starttime}" endtime="{endtime}">msg</status>
      </test>
      <status status="{status}" starttime="{starttime}" endtime="{endtime}"/>
    </suite>
    <status status="{status}" starttime="{starttime}" endtime="{endtime}"/>
  </suite>
  <statistics><total><stat pass="1" fail="0" skip="0">All Tests</stat></total></statistics>
  <errors/>
</robot>
"""
    file_path.write_text(content, encoding="utf-8")


def _write_empty_robot_output(file_path: Path, *, root_name: str) -> None:
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<robot generator="Robot 6.1.1" generated="20260101 10:00:00.000" rpa="false" schemaversion="4">
  <suite id="s1" name="{root_name}" source="{root_name}.robot">
    <status status="PASS" starttime="20260101 10:00:00.000" endtime="20260101 10:00:01.000"/>
  </suite>
  <statistics><total><stat pass="0" fail="0" skip="0">All Tests</stat></total></statistics>
  <errors/>
</robot>
"""
    file_path.write_text(content, encoding="utf-8")


class TestControllerLoading(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        QSettings("Robotframework-XMLMergeTool", "MainWindow").clear()
        self.warning_patcher = patch("xml_merge_tool.ui.controller.QMessageBox.warning")
        self.critical_patcher = patch("xml_merge_tool.ui.controller.QMessageBox.critical")
        self.information_patcher = patch(
            "xml_merge_tool.ui.controller.QMessageBox.information"
        )
        self.mock_warning = self.warning_patcher.start()
        self.mock_critical = self.critical_patcher.start()
        self.mock_information = self.information_patcher.start()
        self.controller = MainController()
        self.controller._run_worker = self._run_worker_sync

    def tearDown(self):
        self.controller.window.close()
        self.warning_patcher.stop()
        self.critical_patcher.stop()
        self.information_patcher.stop()

    def _run_worker_sync(
        self, *, busy_message: str, fn, args, on_success
    ):
        del busy_message
        on_success(fn(*args))
        self.app.processEvents()

    def _selected_candidates(self) -> list[str]:
        return [
            self.controller._candidate_by_id[candidate_id].file_name
            for candidate_id in self.controller.window.get_selected_ids()
        ]

    def test_append_file_keeps_existing_candidates(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            xml_b = root / "b.xml"
            _write_robot_output(xml_a, root_name="RootA", leaf_name="SuiteA")
            _write_robot_output(xml_b, root_name="RootB", leaf_name="SuiteB")

            self.controller.load_files([str(xml_a)])
            self.controller.load_files([str(xml_b)])

            self.assertEqual(
                self.controller._active_file_paths,
                [str(xml_a.resolve()), str(xml_b.resolve())],
            )
            self.assertEqual(
                sorted(candidate.suite_path for candidate in self.controller._candidates),
                ["RootA / SuiteA", "RootB / SuiteB"],
            )
            self.assertEqual(self.controller._parsed_file_count, 2)
            self.assertFalse(self.mock_warning.called)
            self.assertFalse(self.mock_critical.called)

    def test_tree_selected_only_can_uncheck_one_file_without_crashing(self):
        with _workspace_temp_dir() as root:
            xml_file = root / "selected.xml"
            _write_robot_output(
                xml_file,
                root_name="Root",
                leaf_name="SelectedSuite",
            )
            self.controller.load_files([str(xml_file)])

            window = self.controller.window
            tree_index = window.view_mode_combo.findData("tree")
            self.assertGreaterEqual(tree_index, 0)
            window.view_mode_combo.setCurrentIndex(tree_index)
            window.selected_only_checkbox.setChecked(True)
            self.app.processEvents()

            selected_id = window.get_selected_ids()[0]
            checkbox_item = window._checkbox_item_by_candidate_id[selected_id]
            checkbox_item.setCheckState(Qt.CheckState.Unchecked)
            self.app.processEvents()

            self.assertEqual(window.get_selected_ids(), [])
            self.assertEqual(window.get_visible_count(), 0)
            self.assertIn("当前勾选 0 个", window.summary_label.text())

    def test_reloading_same_path_replaces_old_candidates_and_recomputes_defaults(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            xml_b = root / "b.xml"
            _write_robot_output(
                xml_a,
                root_name="Root",
                leaf_name="SuiteA",
                status="PASS",
                endtime="20260101 10:00:01.000",
            )
            _write_robot_output(
                xml_b,
                root_name="Root",
                leaf_name="SuiteA",
                status="SKIP",
                endtime="20260101 10:00:02.000",
            )

            self.controller.load_files([str(xml_a), str(xml_b)])
            self.assertEqual(self._selected_candidates(), ["b.xml"])

            _write_robot_output(
                xml_a,
                root_name="Root",
                leaf_name="SuiteA",
                status="PASS",
                endtime="20260101 10:00:03.000",
            )
            self.controller.load_files([str(xml_a)])

            self.assertEqual(
                self.controller._active_file_paths,
                [str(xml_b.resolve()), str(xml_a.resolve())],
            )
            self.assertEqual(len(self.controller._candidates), 2)
            self.assertEqual(self._selected_candidates(), ["a.xml"])

    def test_failed_reload_of_same_path_removes_stale_candidates(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            _write_robot_output(xml_a, root_name="RootA", leaf_name="SuiteA")

            self.controller.load_files([str(xml_a)])
            self.assertEqual(len(self.controller._candidates), 1)

            xml_a.write_bytes(b"\x00\x01broken-xml")
            self.controller.load_files([str(xml_a)])

            self.assertEqual(self.controller._active_file_paths, [])
            self.assertEqual(len(self.controller._candidates), 0)
            self.assertEqual(self.controller._parsed_file_count, 0)
            self.assertEqual(self.controller._failed_file_count, 1)
            self.mock_critical.assert_called_once()

    def test_zero_candidate_file_stays_in_active_file_set(self):
        with _workspace_temp_dir() as root:
            xml_empty = root / "empty.xml"
            xml_b = root / "b.xml"
            _write_empty_robot_output(xml_empty, root_name="EmptyRoot")
            _write_robot_output(xml_b, root_name="RootB", leaf_name="SuiteB")

            self.controller.load_files([str(xml_empty)])
            self.assertEqual(self.controller._active_file_paths, [str(xml_empty.resolve())])
            self.assertEqual(len(self.controller._candidates), 0)
            self.assertIn("未加载到可用套件", self.controller.window.summary_label.text())

            self.controller.load_files([str(xml_b)])

            self.assertEqual(
                self.controller._active_file_paths,
                [str(xml_empty.resolve()), str(xml_b.resolve())],
            )
            self.assertEqual(self.controller._parsed_file_count, 2)
            self.assertEqual(len(self.controller._candidates), 1)

    def test_loading_files_in_tree_view_keeps_rows_visible(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            xml_b = root / "b.xml"
            _write_robot_output(xml_a, root_name="RootA", leaf_name="SuiteA")
            _write_robot_output(xml_b, root_name="RootB", leaf_name="SuiteB")

            self.controller.window.view_mode_combo.setCurrentText("路径树")
            self.app.processEvents()

            self.controller.load_files([str(xml_a), str(xml_b)])

            self.assertTrue(self.controller.window.is_tree_view_mode())
            self.assertGreater(self.controller.window.tree_view.model().rowCount(), 0)
            self.assertEqual(
                set(self.controller.window.get_filtered_candidate_ids()),
                set(self.controller._candidate_by_id),
            )
            self.assertIn("可见 2 个", self.controller.window.summary_label.text())

    def test_tree_filter_widgets_refresh_visible_candidates_via_connected_signals(self):
        with _workspace_temp_dir() as root:
            xml_a1 = root / "a1.xml"
            xml_a2 = root / "a2.xml"
            xml_b = root / "b.xml"
            _write_robot_output(
                xml_a1,
                root_name="RootA",
                leaf_name="SuiteA",
                status="PASS",
                endtime="20260101 10:00:01.000",
            )
            _write_robot_output(
                xml_a2,
                root_name="RootA",
                leaf_name="SuiteA",
                status="FAIL",
                endtime="20260101 10:00:02.000",
            )
            _write_robot_output(
                xml_b,
                root_name="RootB",
                leaf_name="SuiteB",
                status="SKIP",
                endtime="20260101 10:00:03.000",
            )

            self.controller.window.view_mode_combo.setCurrentText("路径树")
            self.app.processEvents()
            self.controller.load_files([str(xml_a1), str(xml_a2), str(xml_b)])

            candidate_ids = {
                (candidate.file_name, candidate.suite_path): candidate_id
                for candidate_id, candidate in self.controller._candidate_by_id.items()
            }

            self.controller.window.selected_only_checkbox.setChecked(True)
            self.app.processEvents()
            self.assertEqual(
                set(self.controller.window.get_filtered_candidate_ids()),
                {
                    candidate_ids[("a1.xml", "RootA / SuiteA")],
                    candidate_ids[("b.xml", "RootB / SuiteB")],
                },
            )

            self.controller.window.selected_only_checkbox.setChecked(False)
            self.controller.window.conflict_only_checkbox.setChecked(True)
            self.app.processEvents()
            self.assertEqual(
                set(self.controller.window.get_filtered_candidate_ids()),
                {
                    candidate_ids[("a1.xml", "RootA / SuiteA")],
                    candidate_ids[("a2.xml", "RootA / SuiteA")],
                },
            )

            self.controller.window.conflict_only_checkbox.setChecked(False)
            self.controller.window.search_edit.setText("RootA")
            self.app.processEvents()
            self.assertEqual(
                set(self.controller.window.get_filtered_candidate_ids()),
                {
                    candidate_ids[("a1.xml", "RootA / SuiteA")],
                    candidate_ids[("a2.xml", "RootA / SuiteA")],
                },
            )

    def test_export_single_candidate_xml_uses_save_dialog_and_adds_extension(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            _write_robot_output(xml_a, root_name="Root", leaf_name="SuiteA")
            self.controller.load_files([str(xml_a)])
            candidate = self.controller._candidates[0]
            requested_path = root / "chosen-name"

            with patch(
                "xml_merge_tool.ui.controller.QFileDialog.getSaveFileName",
                return_value=(str(requested_path), "XML Files (*.xml)"),
            ) as save_dialog:
                self.controller.export_single_candidate_xml(candidate.candidate_id)

            exported_path = root / "chosen-name.xml"
            self.assertTrue(exported_path.exists())
            default_path = save_dialog.call_args.args[2]
            self.assertEqual(Path(default_path).name, "SuiteA.xml")
            self.mock_information.assert_called_once()
            self.assertIn(str(exported_path), self.mock_information.call_args.args[2])

    def test_export_single_candidate_xml_cancel_is_noop(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            _write_robot_output(xml_a, root_name="Root", leaf_name="SuiteA")
            self.controller.load_files([str(xml_a)])
            candidate = self.controller._candidates[0]

            with patch(
                "xml_merge_tool.ui.controller.QFileDialog.getSaveFileName",
                return_value=("", ""),
            ):
                self.controller.export_single_candidate_xml(candidate.candidate_id)

            self.mock_information.assert_not_called()
            self.assertEqual(list(root.glob("*.xml")), [xml_a])

    def test_export_selected_tests_excel_includes_checked_items_hidden_by_filter(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            xml_b = root / "b.xml"
            _write_robot_output(xml_a, root_name="RootA", leaf_name="SuiteA")
            _write_robot_output(xml_b, root_name="RootB", leaf_name="SuiteB")
            self.controller.load_files([str(xml_a), str(xml_b)])
            self.assertEqual(len(self.controller.window.get_selected_ids()), 2)

            self.controller.window.search_edit.setText("RootA")
            self.app.processEvents()
            self.assertEqual(self.controller.window.get_visible_count(), 1)

            requested_path = root / "checked-tests"
            with patch(
                "xml_merge_tool.ui.controller.QFileDialog.getSaveFileName",
                return_value=(str(requested_path), "Excel Files (*.xlsx)"),
            ) as save_dialog:
                self.controller.export_selected_tests_excel()

            exported_path = root / "checked-tests.xlsx"
            self.assertTrue(exported_path.exists())
            self.assertRegex(
                Path(save_dialog.call_args.args[2]).name,
                r"selected-tests-\d{8}-\d{6}\.xlsx",
            )
            workbook = load_workbook(exported_path)
            try:
                worksheet = workbook["测试结果"]
                self.assertEqual(worksheet.max_row, 3)
                self.assertEqual(
                    [worksheet["A2"].value, worksheet["A3"].value],
                    ["case1", "case1"],
                )
                self.assertEqual(
                    [worksheet["B2"].value, worksheet["B3"].value],
                    ["SuiteA.robot", "SuiteB.robot"],
                )
                self.assertEqual(
                    [worksheet["D2"].value, worksheet["D3"].value],
                    [timedelta(seconds=1), timedelta(seconds=1)],
                )
                self.assertEqual(worksheet["D2"].number_format, "[h]:mm:ss.000")
            finally:
                workbook.close()
            self.mock_information.assert_called_once()
            message = self.mock_information.call_args.args[2]
            self.assertIn("已导出 2 条测试", message)
            self.assertIn(str(exported_path), message)

    def test_export_selected_tests_excel_cancel_and_empty_selection_are_safe(self):
        with _workspace_temp_dir() as root:
            xml_a = root / "a.xml"
            _write_robot_output(xml_a, root_name="Root", leaf_name="SuiteA")
            self.controller.load_files([str(xml_a)])

            with patch(
                "xml_merge_tool.ui.controller.QFileDialog.getSaveFileName",
                return_value=("", ""),
            ):
                self.controller.export_selected_tests_excel()
            self.assertEqual(list(root.glob("*.xlsx")), [])
            self.mock_information.assert_not_called()

            self.controller.window.apply_selection([])
            self.controller.export_selected_tests_excel()
            self.mock_critical.assert_called_once()
            self.assertIn("至少勾选", self.mock_critical.call_args.args[2])

    def test_normalize_xlsx_path_replaces_wrong_extension(self):
        self.assertEqual(
            self.controller._normalize_xlsx_path("D:/tmp/result"),
            "D:/tmp/result.xlsx",
        )
        self.assertEqual(
            self.controller._normalize_xlsx_path("D:/tmp/result.xls"),
            "D:/tmp/result.xlsx",
        )
        self.assertEqual(
            self.controller._normalize_xlsx_path("D:/tmp/result.XLSX"),
            "D:/tmp/result.XLSX",
        )

    def test_safe_export_name_replaces_windows_invalid_names(self):
        self.assertEqual(
            self.controller._safe_export_name('Suite:A/B*?'),
            "Suite_A_B__",
        )
        self.assertEqual(self.controller._safe_export_name("CON"), "_CON")


if __name__ == "__main__":
    unittest.main()
